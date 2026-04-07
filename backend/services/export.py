"""Export RFP analysis to Word (.docx) and PDF."""
from __future__ import annotations

import io
import json
from typing import Any


def _loads(val: Any, default: Any) -> Any:
    if not val:
        return default
    try:
        return json.loads(val)
    except Exception:
        return default


def _safe_date(rfp) -> str:
    try:
        return rfp.created_at.strftime("%B %d, %Y")
    except Exception:
        return ""


# ── Word ──────────────────────────────────────────────────────────────────────

def build_docx(rfp) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()

    # Remove default empty paragraph
    for para in doc.paragraphs:
        p = para._element
        p.getparent().remove(p)

    # Title
    doc.add_heading(rfp.filename, 0)

    # Metadata line
    industry = (rfp.industry or "general").title()
    meta = doc.add_paragraph(f"{_safe_date(rfp)}   ·   Industry: {industry}")
    meta.runs[0].font.size = Pt(10)
    meta.runs[0].font.color.rgb = RGBColor(0x6B, 0x8F, 0x72)
    doc.add_paragraph()

    # ── Score ─────────────────────────────────────────────────────────────────
    doc.add_heading("Score & Decision", 1)
    decision = rfp.decision or "NO BID"
    score = rfp.score or 0
    p = doc.add_paragraph()
    run = p.add_run(f"{decision}  —  {score}/100")
    run.bold = True
    run.font.size = Pt(14)

    breakdown = _loads(rfp.score_breakdown, {})
    if breakdown:
        doc.add_heading("Breakdown", 2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Category"
        hdr[1].text = "Score"
        for k, v in breakdown.items():
            row = table.add_row().cells
            row[0].text = k.replace("_", " ").title()
            row[1].text = str(v)

    if rfp.reasoning:
        doc.add_heading("Reasoning", 2)
        doc.add_paragraph(rfp.reasoning)

    doc.add_paragraph()

    # ── Requirements ─────────────────────────────────────────────────────────
    requirements = _loads(rfp.requirements, {})
    if requirements:
        doc.add_heading("Requirements", 1)
        for category, items in requirements.items():
            doc.add_heading(category.replace("_", " ").title(), 2)
            reqs = items if isinstance(items, list) else []
            for req in reqs:
                if isinstance(req, dict):
                    text = req.get("description") or req.get("text") or str(req)
                    priority = req.get("priority", "")
                    p = doc.add_paragraph(text, style="List Bullet")
                    if priority:
                        p.add_run(f"  [{priority}]").italic = True
                else:
                    doc.add_paragraph(str(req), style="List Bullet")
        doc.add_paragraph()

    # ── Risks ─────────────────────────────────────────────────────────────────
    risks = _loads(rfp.risks, [])
    if risks:
        doc.add_heading("Risks", 1)
        for risk in risks:
            if isinstance(risk, dict):
                title_text = risk.get("title") or risk.get("name") or "Risk"
                severity = risk.get("severity", "")
                description = risk.get("description") or risk.get("mitigation") or ""
                label = f"{title_text} ({severity})" if severity else title_text
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(label).bold = True
                if description:
                    doc.add_paragraph(description, style="List Continue")
            else:
                doc.add_paragraph(str(risk), style="List Bullet")
        doc.add_paragraph()

    # ── Proposal ──────────────────────────────────────────────────────────────
    if rfp.proposal:
        doc.add_heading("Proposal", 1)
        for line in rfp.proposal.split("\n"):
            stripped = line.strip()
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], 3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], 2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:], 2)
            elif stripped.startswith("- ") or stripped.startswith("• "):
                doc.add_paragraph(stripped[2:], style="List Bullet")
            elif stripped:
                doc.add_paragraph(stripped)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────

def build_pdf(rfp) -> bytes:
    from fpdf import FPDF

    class _PDF(FPDF):
        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(150, 150, 150)
            self.cell(0, 10, f"Page {self.page_no()}", align="C")

    pdf = _PDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_margins(20, 20, 20)

    # ── helpers ───────────────────────────────────────────────────────────────
    def section_title(text: str) -> None:
        pdf.ln(5)
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(40, 90, 55)
        pdf.cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(60, 130, 75)
        pdf.set_line_width(0.3)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + pdf.epw, pdf.get_y())
        pdf.ln(3)

    def subsection(text: str) -> None:
        pdf.ln(2)
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(55, 100, 65)
        pdf.cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)

    def body(text: str, indent: float = 0) -> None:
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(30, 30, 30)
        if indent:
            pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(pdf.epw - indent, 5, text)
        pdf.set_x(pdf.l_margin)

    def bullet(text: str) -> None:
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(30, 30, 30)
        pdf.set_x(pdf.l_margin + 4)
        pdf.cell(5, 5, "\u2022")
        pdf.multi_cell(pdf.epw - 9, 5, text)
        pdf.set_x(pdf.l_margin)

    # ── Title ─────────────────────────────────────────────────────────────────
    pdf.set_font("Helvetica", "B", 16)
    pdf.set_text_color(20, 60, 35)
    pdf.multi_cell(0, 9, rfp.filename)
    pdf.ln(1)

    industry = (rfp.industry or "general").title()
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(100, 140, 110)
    pdf.cell(0, 6, f"{_safe_date(rfp)}   ·   Industry: {industry}", new_x="LMARGIN", new_y="NEXT")

    # ── Score ─────────────────────────────────────────────────────────────────
    section_title("Score & Decision")
    decision = rfp.decision or "NO BID"
    score = rfp.score or 0
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(*(40, 120, 60) if decision == "BID" else (160, 40, 40))
    pdf.cell(0, 11, f"{decision}  —  {score}/100", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    breakdown = _loads(rfp.score_breakdown, {})
    if breakdown:
        subsection("Breakdown")
        for k, v in breakdown.items():
            label = k.replace("_", " ").title()
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(50, 50, 50)
            pdf.cell(85, 6, f"  {label}")
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(0, 6, str(v), new_x="LMARGIN", new_y="NEXT")

    if rfp.reasoning:
        subsection("Reasoning")
        body(rfp.reasoning)

    # ── Requirements ─────────────────────────────────────────────────────────
    requirements = _loads(rfp.requirements, {})
    if requirements:
        section_title("Requirements")
        for category, items in requirements.items():
            subsection(category.replace("_", " ").title())
            reqs = items if isinstance(items, list) else []
            for req in reqs:
                if isinstance(req, dict):
                    text = req.get("description") or req.get("text") or str(req)
                    priority = req.get("priority", "")
                    bullet(f"{text}  [{priority}]" if priority else text)
                else:
                    bullet(str(req))

    # ── Risks ─────────────────────────────────────────────────────────────────
    risks = _loads(rfp.risks, [])
    if risks:
        section_title("Risks")
        for risk in risks:
            if isinstance(risk, dict):
                title_text = risk.get("title") or risk.get("name") or "Risk"
                severity = risk.get("severity", "")
                description = risk.get("description") or risk.get("mitigation") or ""
                label = f"{title_text} ({severity})" if severity else title_text
                bullet(label)
                if description:
                    body(description, indent=9)
            else:
                bullet(str(risk))

    # ── Proposal ──────────────────────────────────────────────────────────────
    if rfp.proposal:
        section_title("Proposal")
        for line in rfp.proposal.split("\n"):
            stripped = line.strip()
            if stripped.startswith("### "):
                subsection(stripped[4:])
            elif stripped.startswith("## ") or stripped.startswith("# "):
                subsection(stripped.lstrip("#").strip())
            elif stripped.startswith("- ") or stripped.startswith("• "):
                bullet(stripped[2:])
            elif stripped:
                body(stripped)
            else:
                pdf.ln(2)

    return bytes(pdf.output())
